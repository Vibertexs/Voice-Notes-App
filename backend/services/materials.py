"""Reads text out of uploaded coursework without leaving the machine."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from backend.config import (
    MAX_MATERIAL_EXTRACTED_CHARS,
    MAX_OFFICE_ARCHIVE_FILES,
    MAX_OFFICE_UNCOMPRESSED_BYTES,
    MAX_PDF_PAGES_FOR_EXTRACTION,
)


def normalize_extracted_text(text: str) -> str:
    """Keep locally extracted material compact and safe to include in model context."""
    normalized = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"[ \t]+\n", "\n", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    if len(normalized) > MAX_MATERIAL_EXTRACTED_CHARS:
        return normalized[:MAX_MATERIAL_EXTRACTED_CHARS].rstrip() + "\n\n[Material text truncated locally.]"
    return normalized


def validate_office_archive(file_path: Path) -> None:
    """Reject Office zip bombs before a document library expands them in memory."""
    with zipfile.ZipFile(file_path) as archive:
        members = archive.infolist()
        if len(members) > MAX_OFFICE_ARCHIVE_FILES:
            raise ValueError("Office archive contains too many files")
        if sum(member.file_size for member in members) > MAX_OFFICE_UNCOMPRESSED_BYTES:
            raise ValueError("Office archive expands beyond the local safety limit")


def extract_material_text(file_path: Path, suffix: str) -> tuple[str, str, str]:
    """Extract text locally; a failed extraction never prevents a student keeping their file."""
    try:
        if suffix in {".txt", ".md"}:
            source_text = file_path.read_bytes().decode("utf-8", errors="replace")
        elif suffix == ".pdf":
            reader = PdfReader(file_path)
            if reader.is_encrypted and not reader.decrypt(""):
                return "", "locked", "Unlock this PDF before Class AI can read it."
            parts: list[str] = []
            for page_number, page in enumerate(reader.pages, start=1):
                if page_number > MAX_PDF_PAGES_FOR_EXTRACTION:
                    break
                try:
                    page_text = page.extract_text(extraction_mode="layout") or ""
                except Exception:
                    page_text = page.extract_text() or ""
                if page_text.strip():
                    parts.append(f"Page {page_number}\n{page_text}")
                if sum(len(part) for part in parts) >= MAX_MATERIAL_EXTRACTED_CHARS:
                    break
            source_text = "\n\n".join(parts)
        elif suffix == ".docx":
            validate_office_archive(file_path)
            document = Document(file_path)
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table_number, table in enumerate(document.tables, start=1):
                rows = [
                    " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    for row in table.rows
                ]
                rows = [row for row in rows if row]
                if rows:
                    parts.append(f"Table {table_number}\n" + "\n".join(rows))
            source_text = "\n\n".join(parts)
        elif suffix == ".pptx":
            validate_office_archive(file_path)
            presentation = Presentation(file_path)
            parts = []
            for slide_number, slide in enumerate(presentation.slides, start=1):
                slide_parts: list[str] = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text.strip():
                        slide_parts.append(shape.text)
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if cells:
                                slide_parts.append(" | ".join(cells))
                if slide_parts:
                    parts.append(f"Slide {slide_number}\n" + "\n".join(slide_parts))
            source_text = "\n\n".join(parts)
        elif suffix in {".doc", ".ppt"}:
            return (
                "",
                "needs_conversion",
                "Export this older file as PDF, .docx, or .pptx for Class AI to read it.",
            )
        else:
            return "", "unsupported", "This file type is attached, but Class AI cannot read it yet."
    except Exception:
        return "", "failed", "The file is attached, but its text could not be read locally."

    extracted_text = normalize_extracted_text(source_text)
    if not extracted_text:
        if suffix == ".pdf":
            return "", "needs_ocr", "No selectable text was found. A scanned PDF needs OCR before Class AI can read it."
        return "", "no_text", "No readable text was found in this file."
    if len(extracted_text) >= MAX_MATERIAL_EXTRACTED_CHARS:
        return extracted_text, "ready", "Text was read locally; the saved AI source is capped for performance."
    return extracted_text, "ready", "Text is ready for Class AI on this laptop."
